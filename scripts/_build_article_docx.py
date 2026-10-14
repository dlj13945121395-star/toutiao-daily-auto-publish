#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build a Word document on the Desktop: title + body text interleaved with
3 inline AI-generated images. Source article: NVIDIA收购Hugging Face."""
import json, os, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

SKILL_ROOT = r"C:\Users\didi\.workbuddy\skills\toutiao-weitoutiao-auto-publish"
post = json.load(open(os.path.join(SKILL_ROOT, "article_post.json"), "r", encoding="utf-8"))
TITLE = post["title"].strip()
PARAS = [p.strip() for p in post["body"].split("\n\n") if p.strip()]

# Images live INSIDE the skill dir (an earlier Desktop\inline_imgs copy was cleaned
# up by the user, which silently produced an image-less docx — keep them here).
IMG_DIR = os.path.join(SKILL_ROOT, "images", "article_inline")
IMGS = [
    ("Cinematic_ultra_wide_photograp_2026-09-09T08-11-13.png",
     "（图 1 · 全球开源 AI 社区的神经网络）",
     "Hugging Face 这个名字"),                                 # after introducing HF
    ("Cinematic_ultra_wide_photograp_2026-09-09T08-12-03.png",
     "（图 2 · 算力地主：分发层换了主人）",
     "算力地主"),                                             # after the chip metaphor section
    ("Cinematic_ultra_wide_photograp_2026-09-09T08-12-36.png",
     "（图 3 · 手里有第二张牌：能换地方）",
     "能换地方"),                                             # after the closing "keep a second card" section
]
INSERT_AFTER = None  # resolved dynamically

# Fail loudly if any image is missing — silently producing an image-less docx before
# is how we ended up shipping a text-only file once.
_missing = [f for f, _, _ in IMGS if not os.path.exists(os.path.join(IMG_DIR, f))]
if _missing:
    print("ERROR: 缺少图片文件，已中止（避免生成无图文档）:", _missing)
    sys.exit(1)

OUT = r"C:\Users\didi\Desktop\英伟达收购Hugging Face-带配图版-2026-09-09.docx"


def set_cn_font(run, size=11, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "宋体")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")


doc = Document()

# page margins
for sec in doc.sections:
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

# header note
hdr = doc.add_paragraph()
hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = hdr.add_run("【配图版 · 3 张 AI 生成插图 · 含「AI生成 / WORKBUDDY」水印】")
set_cn_font(r, size=9, color=(0x80, 0x80, 0x80))

# title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(10)
r = p.add_run(TITLE)
set_cn_font(r, size=20, bold=True, color=(0x10, 0x2A, 0x5A))

# body with image interleaving
img_cursor = 0
n_imgs = len(IMGS)
for idx, para_text in enumerate(PARAS):
    # add paragraph
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(para_text)
    set_cn_font(r, size=11)

    # insert image if THIS section contains the marker for the next image
    while img_cursor < n_imgs and IMGS[img_cursor][2] in para_text:
        fname, caption, _ = IMGS[img_cursor]
        img_path = os.path.join(IMG_DIR, fname)
        img_cursor += 1
        # caption
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(2)
        cap.paragraph_format.space_after = Pt(10)
        rcap = cap.add_run(caption)
        set_cn_font(rcap, size=9, color=(0x60, 0x60, 0x60))
        # image
        ip = doc.add_paragraph()
        ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ip.paragraph_format.space_after = Pt(8)
        run = ip.add_run()
        try:
            run.add_picture(img_path, width=Cm(14.5))
        except Exception as e:
            print(f"!! 插入图片失败 {img_path}: {e}", flush=True)

# append any remaining images that didn't match (e.g. marker text was missing)
while img_cursor < n_imgs:
    fname, caption, _ = IMGS[img_cursor]
    img_path = os.path.join(IMG_DIR, fname)
    img_cursor += 1
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    set_cn_font(cap.add_run(caption), size=9, color=(0x60, 0x60, 0x60))
    ip = doc.add_paragraph()
    ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = ip.add_run()
    try:
        run.add_picture(img_path, width=Cm(14.5))
    except Exception as e:
        print(f"!! 插入图片失败 {img_path}: {e}", flush=True)

# footer
foot = doc.add_paragraph()
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
foot.paragraph_format.space_before = Pt(8)
r = foot.add_run("— 完 —")
set_cn_font(r, size=10, color=(0x80, 0x80, 0x80))

doc.save(OUT)
print(f"OK 已生成: {OUT}")
print(f"   段落 {len(PARAS)} 段，图片 {img_cursor} 张。")